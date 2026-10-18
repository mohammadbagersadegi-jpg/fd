#!/usr/bin/env python3
"""Fill the existing four-page DOCX form with the simple 1 Aban plan.

The source Word file is used only as a visual template.  The generator replaces
all twelve daily tables, including the formerly empty sixth task cell, and
keeps the existing page geometry, branding, embedded assets, and RTL styles.

Dependencies:
    pip install -r requirements-docx.txt
"""

from __future__ import annotations

import argparse
import json
import zipfile
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE = (
    ROOT / "برنامه جدیید" / "برنامه-اصلاح‌شده-رتبه‌لند-۱-آبان-۱۴۰۵.docx"
)
DEFAULT_DATA = ROOT / "examples" / "1aban-1405-simple-formula.json"
DEFAULT_OUTPUT = ROOT / "برنامه جدیید" / "برنامه-فرمول-جدید-۱-آبان-۱۴۰۵.docx"

DAY_DATES = [
    "۱۸ مهر",
    "۱۹ مهر",
    "۲۰ مهر",
    "۲۱ مهر",
    "۲۲ مهر",
    "۲۳ مهر",
    "۲۵ مهر",
    "۲۶ مهر",
    "۲۷ مهر",
    "۲۸ مهر",
    "۲۹ مهر",
    "۳۰ مهر",
]
PAGE_HEADER_TABLES = [0, 6, 13, 19]
PAGE_SUMMARY_TABLES = [1, 7, 14, 20]
DAY_TABLES = [2, 3, 4, 8, 9, 10, 15, 16, 17, 21, 22, 23]
FRIDAY_TABLES = [11, 24]

PERSIAN_DIGITS = str.maketrans("0123456789", "۰۱۲۳۴۵۶۷۸۹")
PROHIBITED_WORDS = ("لقمه", "موزاییک", "قاچ")


def fa(value: object) -> str:
    return str(value).translate(PERSIAN_DIGITS)


def duration_text(minutes: int) -> str:
    if minutes == 0:
        return "—"
    hours, rest = divmod(minutes, 60)
    if rest == 0:
        return fa(hours)
    return f"{fa(hours)}:{fa(rest):0>2}"


def total_text(minutes: int) -> str:
    hours, rest = divmod(minutes, 60)
    return f"{fa(hours)}:{fa(rest):0>2}"


def recursive_text(table) -> str:
    return "".join(node.text or "" for node in table._tbl.iter() if node.tag.endswith("}t"))


def set_paragraph_text(paragraph, value: str) -> None:
    """Replace text while retaining the paragraph and first-run formatting."""
    text_nodes = paragraph._p.xpath(".//w:t")
    if not text_nodes:
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        run.append(text)
        paragraph._p.append(run)
        text_nodes = [text]
    text_nodes[0].text = value
    for node in text_nodes[1:]:
        node.text = ""


def set_cell_text(cell, value: str) -> None:
    if not cell.paragraphs:
        cell.add_paragraph()
    set_paragraph_text(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        set_paragraph_text(paragraph, "")


def replace_task_cell(cell, paragraph_templates: list, box: dict) -> None:
    tc = cell._tc
    tc_pr = tc.tcPr
    for child in list(tc):
        if child is not tc_pr:
            tc.remove(child)
    for paragraph_xml in paragraph_templates:
        tc.append(deepcopy(paragraph_xml))

    paragraphs = cell.paragraphs
    if len(paragraphs) != 3:
        raise ValueError(f"Expected three task paragraphs, found {len(paragraphs)}")
    set_paragraph_text(paragraphs[0], f"☐ {box['subject']}")
    set_paragraph_text(paragraphs[1], box["topic"])
    set_paragraph_text(paragraphs[2], box["detail"])


def set_metric(metric_cell, minutes: int, tests: int) -> None:
    if len(metric_cell.tables) != 1:
        raise ValueError("A task metric cell must contain one nested table")
    row = metric_cell.tables[0].rows[0]
    if len(row.cells) != 5:
        raise ValueError("Unexpected task metric geometry")
    set_cell_text(row.cells[1], duration_text(minutes))
    set_cell_text(row.cells[4], fa(tests) if tests else "۰")


def validate_data(data: dict) -> None:
    days = data.get("days", [])
    if len(days) != 12:
        raise ValueError(f"Expected 12 study days, found {len(days)}")
    if [day["date"] for day in days] != DAY_DATES:
        raise ValueError("Unexpected dates or ordering in plan data")

    expected_tests = {1: 265, 2: 380}
    for week in (1, 2):
        current = [day for day in days if day["week"] == week]
        if len(current) != 6:
            raise ValueError(f"Week {week} must contain six days")
        if any(len(day["boxes"]) != 6 for day in current):
            raise ValueError(f"Week {week}: every day must contain six form cells")

        minutes = sum(box["minutes"] for day in current for box in day["boxes"])
        tests = sum(box["tests"] for day in current for box in day["boxes"])
        sessions = sum(
            box["minutes"] == data["meta"]["session_minutes"]
            for day in current
            for box in day["boxes"]
        )
        finals = sum(
            box["minutes"]
            for day in current
            for box in day["boxes"]
            if box["subject"] == "امتحانات نهایی"
        )
        if minutes != 2715:
            raise ValueError(f"Week {week}: expected 2715 planned minutes, got {minutes}")
        if sessions != 33:
            raise ValueError(f"Week {week}: expected 33 specialized sessions, got {sessions}")
        if finals != 240:
            raise ValueError(f"Week {week}: expected four final-exam hours, got {finals}")
        if tests != expected_tests[week]:
            raise ValueError(
                f"Week {week}: expected {expected_tests[week]} tests, got {tests}"
            )

    full_text = json.dumps(data, ensure_ascii=False)
    for word in PROHIBITED_WORDS:
        if word in full_text:
            raise ValueError(f"Specialized jargon is not allowed in this output: {word}")
    if any(box["minutes"] not in {0, 75, 120} for day in days for box in day["boxes"]):
        raise ValueError("Only 0, 75, and 120 minute task cells are supported")
    if any(day["tests"] if "tests" in day else False for day in days):
        raise ValueError("Daily totals must be computed, not stored")


def locate_day_table(document: Document, date: str):
    matches = [
        table
        for table in document.tables
        if f"| {date}" in recursive_text(table) and len(table.rows) == 5
    ]
    if len(matches) != 1:
        raise ValueError(f"Expected one daily table for {date}, found {len(matches)}")
    return matches[0]


def populate_page_headers(document: Document, data: dict) -> None:
    week_tests = {
        week: sum(
            box["tests"]
            for day in data["days"]
            if day["week"] == week
            for box in day["boxes"]
        )
        for week in (1, 2)
    }

    for page_number, table_index in enumerate(PAGE_HEADER_TABLES):
        table = document.tables[table_index]
        title_table = table.rows[0].cells[1].tables[0]
        set_cell_text(title_table.rows[0].cells[0], "برنامه با فرمول جدید")

        info_table = table.rows[0].cells[2].tables[0]
        week = 1 if page_number < 2 else 2
        week_label = "هفته: ۱ ـ دوازدهم" if week == 1 else "هفته: ۲ ـ مرور پایهٔ دهم"
        set_cell_text(info_table.rows[0].cells[0], week_label)

    for page_number, table_index in enumerate(PAGE_SUMMARY_TABLES):
        week = 1 if page_number < 2 else 2
        table = document.tables[table_index]
        hour_table = table.rows[0].cells[3].tables[0]
        test_table = table.rows[0].cells[4].tables[0]
        set_cell_text(hour_table.rows[0].cells[1], "۴۵:۱۵*")
        set_cell_text(test_table.rows[0].cells[1], fa(week_tests[week]))


def populate_day(document: Document, day: dict) -> None:
    table = locate_day_table(document, day["date"])
    if [len(row.cells) for row in table.rows] != [3, 3, 3, 3, 3]:
        raise ValueError(f"{day['date']}: unexpected daily table geometry")

    header = table.rows[0].cells[0].tables[0].rows[0]
    set_cell_text(header.cells[1], f"اهداف: {day['objective']}")
    minutes = sum(box["minutes"] for box in day["boxes"])
    tests = sum(box["tests"] for box in day["boxes"])
    set_cell_text(header.cells[3], total_text(minutes))
    set_cell_text(header.cells[6], fa(tests))

    first_row_template = [deepcopy(p._p) for p in table.rows[1].cells[0].paragraphs]
    second_row_template = [deepcopy(p._p) for p in table.rows[3].cells[0].paragraphs]
    if len(first_row_template) != 3 or len(second_row_template) != 3:
        raise ValueError(f"{day['date']}: source task cell is not a three-line cell")

    task_cells = list(table.rows[1].cells) + list(table.rows[3].cells)
    metric_cells = list(table.rows[2].cells) + list(table.rows[4].cells)
    for index, (task_cell, metric_cell, box) in enumerate(
        zip(task_cells, metric_cells, day["boxes"], strict=True)
    ):
        templates = first_row_template if index < 3 else second_row_template
        replace_task_cell(task_cell, templates, box)
        set_metric(metric_cell, box["minutes"], box["tests"])


def populate_friday(document: Document, data: dict) -> None:
    specs = [
        (FRIDAY_TABLES[0], data["fridays"]["week1"]),
        (FRIDAY_TABLES[1], data["fridays"]["target"]),
    ]
    for table_index, spec in specs:
        table = document.tables[table_index]
        header = table.rows[0].cells[0].tables[0].rows[0]
        set_cell_text(header.cells[1], spec["title"])
        task_cell = table.rows[1].cells[0]
        if len(task_cell.paragraphs) < 2:
            raise ValueError("Friday task cell is missing its detail paragraph")
        set_paragraph_text(task_cell.paragraphs[0], "☐ توضیح برنامه:")
        set_paragraph_text(task_cell.paragraphs[1], spec["detail"])


def verify_output(output: Path, data: dict) -> dict[str, object]:
    if not output.is_file():
        raise FileNotFoundError(output)
    with zipfile.ZipFile(output) as package:
        corrupt = package.testzip()
        if corrupt:
            raise ValueError(f"Corrupt DOCX member: {corrupt}")
        names = package.namelist()
        document_xml = package.read("word/document.xml").decode("utf-8")
        app_xml = package.read("docProps/app.xml")

    if b"<Pages>4</Pages>" not in app_xml:
        raise ValueError("The stored four-page template metadata changed")
    if document_xml.count('w:type="page"') != 3:
        raise ValueError("Expected three explicit page breaks")

    result = Document(output)
    all_text = "\n".join(
        node.text or "" for node in result.element.body.iter() if node.tag.endswith("}t")
    )
    for word in PROHIBITED_WORDS:
        if word in all_text:
            raise ValueError(f"Generated Word file still contains jargon: {word}")
    for day in data["days"]:
        for box in day["boxes"]:
            for value in (box["subject"], box["topic"], box["detail"]):
                if value not in all_text:
                    raise ValueError(f"Generated Word file is missing: {value}")

    week_tests = {
        week: sum(
            box["tests"]
            for day in data["days"]
            if day["week"] == week
            for box in day["boxes"]
        )
        for week in (1, 2)
    }
    if all_text.count("۴۵:۱۵*") != 4:
        raise ValueError("Weekly planned-time target was not written on all pages")
    if all_text.count(fa(week_tests[1])) < 2 or all_text.count(fa(week_tests[2])) < 2:
        raise ValueError("Weekly test totals are missing")

    return {
        "pages": 4,
        "days": len(data["days"]),
        "form_cells": sum(len(day["boxes"]) for day in data["days"]),
        "planned_minutes": sum(
            box["minutes"] for day in data["days"] for box in day["boxes"]
        ),
        "tests": sum(box["tests"] for day in data["days"] for box in day["boxes"]),
        "package_parts": len(names),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE)
    parser.add_argument("--data", type=Path, default=DEFAULT_DATA)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.template.is_file():
        raise FileNotFoundError(f"Template not found: {args.template}")
    data = json.loads(args.data.read_text(encoding="utf-8"))
    validate_data(data)

    document = Document(args.template)
    if len(document.tables) != 26:
        raise ValueError(f"Expected 26 top-level tables, found {len(document.tables)}")
    if PAGE_HEADER_TABLES != [0, 6, 13, 19] or DAY_TABLES != [2, 3, 4, 8, 9, 10, 15, 16, 17, 21, 22, 23]:
        raise AssertionError("Template map was changed without updating validation")

    populate_page_headers(document, data)
    for day in data["days"]:
        populate_day(document, day)
    populate_friday(document, data)

    document.core_properties.title = data["meta"]["title"]
    document.core_properties.subject = "برنامهٔ آزمون ۱ آبان با فرمول سادهٔ ۷۵دقیقه‌ای"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.output)

    report = verify_output(args.output, data)
    print(json.dumps(report, ensure_ascii=False, indent=2))
    print(args.output)


if __name__ == "__main__":
    main()
